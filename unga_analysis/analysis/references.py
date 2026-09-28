"""Read supplied background documents and date-stamped official UN context."""
from pathlib import Path
from html.parser import HTMLParser
from urllib.request import Request, urlopen
import hashlib
import re

from .common import obj,arr,STR,SYSTEM,ask,quote_in,save,now
from ..io import digest

OFFICIAL_URLS = [
    'https://www.un.org/independent-international-scientific-panel-ai/en/faq',
    'https://www.un.org/global-dialogue-ai-governance/en/faq',
    'https://www.un.org/digital-emerging-technologies/ai-advisory-body',
]


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__();self.skip=0;self.parts=[]
    def handle_starttag(self,tag,attrs):
        if tag in ('script','style'):self.skip+=1
    def handle_endtag(self,tag):
        if tag in ('script','style'):self.skip=max(0,self.skip-1)
    def handle_data(self,data):
        if not self.skip and data.strip():self.parts.append(data.strip())


def reference_blocks(root,out,fetch_current=True):
    blocks=[];inventory=[]
    for index,path in enumerate(sorted((Path(root)/'reference').glob('*')),1):
        if path.suffix.lower() not in ('.pdf','.docx','.txt','.md'):continue
        role='draft' if 'draft' in path.name.lower() else 'supplied_background'
        entry=dict(file=path.relative_to(root).as_posix(),sha256=digest(path),role=role)
        extracted=[]
        if path.suffix.lower()=='.pdf':
            import fitz
            with fitz.open(path) as document:
                extracted=[(dict(pdf_page=i+1),p.get_text()) for i,p in enumerate(document)]
        elif path.suffix.lower()=='.docx':
            from docx import Document
            document=Document(path)
            extracted=[(dict(paragraph=i+1),p.text) for i,p in enumerate(document.paragraphs) if p.text.strip()]
            for ti,t in enumerate(document.tables):
                extracted.extend((dict(table=ti+1,row=ri+1),' | '.join(c.text for c in row.cells)) for ri,row in enumerate(t.rows))
        else:extracted=[(dict(line=i+1),t) for i,t in enumerate(path.read_text(encoding='utf-8').splitlines()) if t.strip()]
        for j,(locator,text) in enumerate(extracted):
            if text.strip():blocks.append(dict(reference_id=f'R{index:02d}_{j+1:04d}',**entry,locator=locator,text=text))
        inventory.append(dict(entry,extracted_blocks=len(extracted)))
    if fetch_current:
        for i,url in enumerate(OFFICIAL_URLS,1):
            try:
                with urlopen(Request(url,headers={'User-Agent':'UNGA-research/1.0'}),timeout=20) as response:
                    raw=response.read(3000000)
                parser=TextHTML();parser.feed(raw.decode('utf-8',errors='replace'))
                text='\n'.join(parser.parts)
                if len(text)<500 or 'Access Denied' in text[:200]:raise ValueError('Unreadable official page')
                content_hash=hashlib.sha256(raw).hexdigest()
                (out/'references').mkdir(parents=True,exist_ok=True)
                (out/'references'/f'official_{i}.html').write_bytes(raw)
                blocks.append(dict(reference_id=f'UN{i}',file=url,sha256=content_hash,role='official_page_retrieved',
                                   locator={'retrieved_at':now()},text=text))
                inventory.append(dict(file=url,sha256=content_hash,status='retrieved',retrieved_at=now()))
            except Exception as exc:
                inventory.append(dict(file=url,status='not_verified_at_run',reason=type(exc).__name__,retrieved_at=now()))
    if out is not None:
        save(out/'reference_inventory.json',inventory)
        save(out/'reference_blocks.json',blocks)
    return blocks,inventory


def summarize_references(blocks,provider,out):
    contract=obj(facts=arr(obj(reference_id=STR,quote=STR,summary=STR,document_status=STR)))
    groups=[];current=[];length=0
    for b in blocks:
        if current and length+len(b['text'])>22000:groups.append(current);current=[];length=0
        current.append(b);length+=len(b['text'])
    if current:groups.append(current)
    facts=[]
    for group in groups:
        source={b['reference_id']:b['text'] for b in group}
        def validate(result):
            for f in result['facts']:
                if f['reference_id'] not in source or not quote_in(f['quote'],source[f['reference_id']]):
                    raise ValueError('Background fact requires an exact supplied quote')
        value=ask(provider,'reference_context',{'blocks':[{k:v for k,v in b.items() if k!='locator'} for b in group]},contract,
            SYSTEM+'Read all supplied background blocks. Summarize at most six relevant facts on institutional mandate/status, chronology, '
            'developing-country capacity needs or SPMU framing. Distinguish drafts, proposals, preliminary scientific reports, adopted decisions '
            'and dated website statements. Older SPMU country counts are historical style context, never current speech evidence. '
            'No supplied or retrieved document proves that an unspecified later event happened. Each fact needs an exact short source quote.',validate)
        facts.extend(value['facts'])
    save(out/'reference_context.json',facts)
    return facts

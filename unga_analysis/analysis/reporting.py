"""Grounded report drafting and flowing four-to-six-page Word/PDF publication."""
from pathlib import Path
from xml.sax.saxutils import escape
import re
import os
import subprocess

from .common import obj,arr,STR,BOOL,SYSTEM,ask,save,table,now,quote_in

TITLES=['Key Facts and Figures','Themes','Role of UN','Implications']
PARAGRAPH=obj(text=STR,evidence_ids=arr(STR),reference_ids=arr(STR),metric_keys=arr(STR))
REPORT=obj(sections=arr(obj(title={'type':'string','enum':TITLES},paragraphs=arr(PARAGRAPH))),
           recommendations=arr(obj(addressee=STR,action=STR,rationale=STR,evidence_ids=arr(STR),metric_keys=arr(STR))))
SECTION_LIMITS={'Key Facts and Figures':(3,110),'Themes':(6,110),'Role of UN':(4,110),'Implications':(3,100)}


def metric_register(stats):
    metrics={}
    for name in ('annual','regional','theme_prevalence','cooccurrence','institution_stance_counts','keywords','matched_panel'):
        for i,row in enumerate(stats[name],1):metrics[f'{name}:{i}']=row
    return metrics


def draft_report(stats,evidence,institutions,taxonomy,reference_facts,provider,out,layout_feedback=None):
    metrics=metric_register(stats)
    source={r['evidence_id']:r for r in evidence}
    refs={r['reference_id'] for r in reference_facts}
    # Bound narrative input while all numeric tables and classifications remain exported.
    year=max(int(y) for y in stats['coverage'])
    latest=[r for r in evidence if r['year']==year]
    historical=[r for r in evidence if r['year']!=year]
    selected=[]
    for code in taxonomy['codes']:
        selected.extend([r for r in latest if r['themes'].get(code)=='Yes'][:4])
        selected.extend([r for r in historical if r['themes'].get(code)=='Yes'][:2])
    selected.extend(latest[:12])
    institutional_ids={r['evidence_id'] for r in institutions}
    selected.extend(r for r in evidence if r['evidence_id'] in institutional_ids and r['year']==year)
    selected=list({r['evidence_id']:r for r in selected}.values())
    selected=selected[:160]
    prompt_metrics={k:v for k,v in metrics.items() if k.startswith('annual:') or
        (k.startswith(('regional:','theme_prevalence:','institution_stance_counts:')) and v['year']==year)}
    for name in ('cooccurrence','keywords','matched_panel'):
        relevant=[(k,v) for k,v in metrics.items() if k.startswith(name+':') and (v.get('year')==year or v.get('year_to')==year)]
        relevant.sort(key=lambda kv:kv[1].get('n_both',kv[1].get('n',0)),reverse=True)
        prompt_metrics.update(relevant[:12])
    selected_ids={r['evidence_id'] for r in selected}
    payload=dict(report_year=year,coverage=stats['coverage'],metrics=prompt_metrics,taxonomy=taxonomy,
                 speech_evidence=[{k:r[k] for k in ('evidence_id','year','iso3','quote','themes','theme_evidence','review_status')} for r in selected],
                 institutions=[{k:r[k] for k in ('evidence_id','year','iso3','mechanism','stances','requested_functions','quote','review_status')} for r in institutions if r['evidence_id'] in selected_ids],
                 background=reference_facts,
                 review_method=stats['review_method'],unresolved_ai_speeches=stats['unresolved_ai_speeches'],
                 incomplete_theme_speeches=stats['incomplete_theme_speeches'])
    def validate(value):
        if [s['title'] for s in value['sections']]!=TITLES:raise ValueError('Exactly four ordered report sections required')
        for section in value['sections']:
            count,words=SECTION_LIMITS[section['title']]
            if not 1<=len(section['paragraphs'])<=count:raise ValueError('Section exceeds its paragraph budget')
            if section['title']=='Themes' and len(section['paragraphs'])<min(4,len(taxonomy['codes'])):raise ValueError('Themes needs a paragraph for each of at least four themes, or every theme if fewer exist')
            for p in section['paragraphs']:
                if len(p['text'].split())>words:raise ValueError('Report paragraph exceeds its section word budget')
                if not (p['evidence_ids'] or p['reference_ids'] or p['metric_keys']):raise ValueError('Each analytical claim needs source or metric references')
                if set(p['evidence_ids'])-selected_ids or set(p['reference_ids'])-refs or set(p['metric_keys'])-set(prompt_metrics):
                    raise ValueError('Unknown report citation')
                if re.search(r'\bE\d{5}\b|\b(?:annual|regional|theme_prevalence):\d+',p['text']):raise ValueError('Keep internal citation keys out of prose; use structured citation fields')
                spans=[]
                for eid in p['evidence_ids']:
                    row=source[eid];spans.append(row['quote']);spans.extend(row.get('theme_evidence',{}).values())
                    spans.extend(r['quote'] for r in row.get('institutions',[]))
                spans.extend(r['quote'] for r in reference_facts if r['reference_id'] in p['reference_ids'])
                for quote in re.findall(r'["“]([^"“”]+)["”]',p['text']):
                    if not any(quote_in(quote,span) for span in spans):raise ValueError('Report quotation must match a cited source span; use unquoted paraphrase for analytical labels')
        if sum(len(p['text'].split()) for s in value['sections'] for p in s['paragraphs'])>1550:raise ValueError('Main narrative exceeds 1550 words; prioritize findings within section budgets')
        if not 3<=len(value['recommendations'])<=5:raise ValueError('Need 3-5 concise recommendations')
        for r in value['recommendations']:
            if len((r['action']+' '+r['rationale']).split())>65:raise ValueError('Recommendation too long')
            if not (r['evidence_ids'] or r['metric_keys']) or set(r['evidence_ids'])-selected_ids or set(r['metric_keys'])-set(prompt_metrics):raise ValueError('Recommendation needs valid supplied speech or metric evidence')
    instructions=SYSTEM+('Draft an English UN EOSG/SPMU strategic report for the latest year, with the four exact sections in order. '
        'Use 2-3 Key Facts paragraphs, 4-6 Themes paragraphs (one per important theme, fewer only if fewer themes exist), '
        '3-4 Role of UN paragraphs and 2-3 Implications paragraphs. Each paragraph at most 110 words (Implications 100). '
        'Use 1200-1550 main narrative words when supported by evidence, and 3-5 recommendations of at most 65 words each. Never pad sparse findings. '
        'Each theme paragraph must explain the finding, country/regional differences and a representative quotation or precise paraphrase. '
        'naming UN/SPMU, Tech Envoy, Panel or Dialogue as addressees. Numbers are already supplied and must not be recomputed or invented. '
        'Tables/charts will be inserted separately. Cite provided evidence IDs, metric keys and background reference IDs. '
        'Every phrase inside quotation marks must match a cited source quote exactly; use unquoted paraphrases for analytical labels. '
        'Distinguish institutional background, Member State positions and your analytical proposals. '
        'Do not imply endorsement, causation, alignment, treaty mandates or completed commitments without evidence. '
        'Treat theme percentages with N<20 as unavailable; discuss counts instead. Partial years are interim, not completed annual trends. '
        'When annual measure is verified_detection_lower_bound, the percentage is verified detections divided by obtained speeches, a minimum, not estimated population prevalence. '
        'Unreviewed speeches are unknown, not No. Screening-selected observations are not a representative sample; never infer absence or adjusted population rates from them. '
        'Theme N is code-specific: AI-positive country-years with resolved 1/0 for that code. Co-mention N requires both codes resolved; matched panels are code-specific. Never assume equal denominators across themes. '
        'Account for automated-review limits and official-versus-ASR changes. Preserve any disagreement or draft status in background sources. '
        'Do not claim current institutional status from undated or draft documents. No rhetorical filler.')
    report=ask(provider,'report_draft',payload,REPORT,instructions,validate,max_tokens=7000)
    # Check factual narrative against the actual source register; allow a corrected full report.
    report=ask(provider,'report_fact_check',dict(payload,draft=report),REPORT,SYSTEM+
        'Audit and correct every claim, number, attribution, institutional status and recommendation in the draft. '
        'Use only the supplied tables and source quotes. Remove unsupported assertions. Keep exactly the specified four sections, '
        'Respect section budgets: Key Facts 1-3 paragraphs, Themes 1-6 (at least four if four themes exist), Role of UN 1-4, Implications 1-3; '
        '110 words per paragraph except Implications 100; at most 1550 narrative words; 3-5 <=65-word recommendations. '
        'Retain theme-specific findings, country/region differences and quotations/paraphrases, with existing citations. '
        'Do not label these two automated passes as human verification.',validate,max_tokens=7000)
    if layout_feedback:
        report=ask(provider,'report_layout_revision',dict(payload,draft=report,layout_feedback=layout_feedback),REPORT,SYSTEM+
            'Revise the draft to fit the measured 4-6 page layout. Aim for 1000-1200 narrative words; preserve four sections, '
            'at least four theme-specific paragraphs if available, country differences, representative quotes/paraphrases and all factual qualifications. '
            'Use up to 90 words per paragraph, concise recommendations and valid existing citations. Never invent or change statistics to save space.',
            validate,max_tokens=7000)
    save(out/'report_content.json',report)
    save(out/'report_metrics.json',metrics)
    keys=list(dict.fromkeys(k for s in report['sections'] for p in s['paragraphs'] for k in p['evidence_ids']+p['reference_ids']+p['metric_keys']))
    keys+=list(k for r in report['recommendations'] for k in r['evidence_ids']+r['metric_keys'] if k not in keys)
    keys=list(dict.fromkeys(keys));background={r['reference_id']:r for r in reference_facts}
    if (out/'reference_blocks.json').exists():
        import json
        blocks={r['reference_id']:r for r in json.loads((out/'reference_blocks.json').read_text(encoding='utf-8'))}
        background={key:dict({k:v for k,v in blocks.get(key,{}).items() if k!='text'},**value) for key,value in background.items()}
    report['citations']={k:dict(number=i+1,source=source.get(k) or background.get(k) or metrics[k]) for i,k in enumerate(keys)}
    save(out/'report_citations.json',report['citations'])
    return report


def charts(stats,destination):
    import matplotlib
    matplotlib.use('Agg')
    from matplotlib import pyplot as plt
    year=max(int(y) for y in stats['coverage'])
    rows=stats['annual']
    fig,ax=plt.subplots(figsize=(7.2,2.35),layout='constrained')
    ax.plot([r['year'] for r in rows],[r['pct'] if r['pct'] is not None else float('nan') for r in rows],color='#174b68',marker='o')
    for r in rows:
        if r['partial'] and r['pct'] is not None:ax.scatter(r['year'],r['pct'],marker='s',s=90,facecolors='white',edgecolors='#ae5c21',zorder=3)
    ax.set(xlabel='General Debate year',ylabel='Verified detections / obtained (%) — minimum' if stats.get('annual_measure')=='verified_detection_lower_bound' else 'AI mention share (%)',ylim=(0,100),xticks=[r['year'] for r in rows])
    ax.grid(axis='y',alpha=.2);ax.spines[['top','right']].set_visible(False)
    trend=destination/'annual_trend.png';fig.savefig(trend,dpi=180);plt.close(fig)
    themes=sorted([r for r in stats['theme_prevalence'] if r['year']==year],key=lambda r:r['n'],reverse=True)[:8]
    fig,ax=plt.subplots(figsize=(7.2,2.7),layout='constrained')
    percent=bool(themes and all(r['pct'] is not None for r in themes))
    if themes:
        ax.barh([f"{r['theme']} ({r['n']}/{r['N']})" for r in themes][::-1],[r['pct'] if percent else r['n'] for r in themes][::-1],color='#247e87')
    else:ax.text(.5,.5,'No resolved theme observations',ha='center',va='center',transform=ax.transAxes)
    ax.set_xlabel('Share of AI-positive countries resolved for each code (%)' if percent else 'Country count; code-specific n/N shown')
    ax.spines[['top','right']].set_visible(False)
    theme=destination/'theme_prevalence.png';fig.savefig(theme,dpi=180);plt.close(fig)
    return trend,theme


def pages(report,stats,institutions,reference_facts,reference_inventory,figures):
    year=max(int(y) for y in stats['coverage']);current=next(r for r in stats['annual'] if r['year']==year)
    partial=current['partial'];sections={s['title']:s for s in report['sections']}
    citations=report.get('citations',{})
    def cite(keys):
        labels=[]
        for key in keys:
            item=citations.get(key)
            if not item:continue
            src=item['source'];label=f"{src['iso3']} {src['year']}" if 'iso3' in src else str(src.get('year','Background'))
            labels.append(f"{label} [{item['number']}]")
        return ' ('+'; '.join(labels)+')' if labels else ''
    def prose(title):
        values=[]
        for p in sections[title]['paragraphs']:
            cites=p['evidence_ids']+p['reference_ids']+p['metric_keys']
            values.append(('paragraph',p['text']+cite(cites)))
        return values
    def ratio(row):return f"{row['n']}/{row['N']} ({row['pct']:.1f}%)" if row['pct'] is not None else f"{row['n']}/{row['N']} (NA)"
    first=[('title',TITLES[0]),('paragraph',f"UNGA{year-1945} | {'INTERIM - partial session' if partial else 'Available session records'} | Analysis date {now()[:10]}. "
        f"Verified AI mentions: {ratio(current)}; {current['obtained']} obtained and {current['uncertain']} unresolved. "+
        ('This is a minimum detection share of obtained addresses, not population prevalence. ' if current.get('measure')=='verified_detection_lower_bound' else 'Denominator: AI-resolved addresses. ')+
        'Unobserved country-years are not counted as negatives.'),('image',str(figures[0]))]
    rows=[['UN regional group','AI mentions n/N','Uncertain']]
    for r in stats['regional']:
        if r['year']==year:rows.append([r['region'],ratio(r),str(r['uncertain'])])
    first.append(('table',rows));first+=prose(TITLES[0])
    first.append(('small','Source changes: official English records (2017-2024); user-accepted automatic transcripts (2025-2026). Squares mark partial years. Regional mapping fixed at 25 September 2026.'))
    second=[('title',TITLES[1]),('image',str(figures[1]))]+prose(TITLES[1])
    top=sorted([r for r in stats['cooccurrence'] if r['year']==year and r['n_both']],key=lambda r:r['n_both'],reverse=True)[:2]
    theme_names={r['code']:r['theme'] for r in stats['theme_prevalence']}
    if top:second.append(('table',[['Theme pair','Joint n / N','P(B|A) / P(A|B)']]+[[theme_names.get(r['theme_a'],r['theme_a'])+' + '+theme_names.get(r['theme_b'],r['theme_b']),f"{r['n_both']} / {r['N']}",f"{r['p_b_given_a']:.2f} / {r['p_a_given_b']:.2f}"] for r in top]))
    terms=sorted([r for r in stats['keywords'] if r['year']==year],key=lambda r:r['n'],reverse=True)[:4]
    if terms:second.append(('table',[['Reviewed expression','Countries n/N','First observed']]+[[r['term'],f"{r['n']}/{r['N']}",str(r['first_observed_year'])] for r in terms]))
    second.append(('small','Theme N: AI-positive country-years resolved for that code; denominators differ by code. Co-mention N requires both codes resolved. Matched-country panels are code-specific. Unknowns remain NA. Multiple labels may sum above 100%. Co-mention is not alignment or a sentence-level link. Full tables include denominators and baselines.'))
    third=[('title',TITLES[2])]+prose(TITLES[2])
    stances=[r for r in stats['institution_stance_counts'] if r['year']==year]
    if stances:third.append(('table',[['Mechanism / stance','n','Countries']]+[[r['mechanism']+' / '+r['stance'],str(r['n']),', '.join(r['countries'][:10])+(f" +{len(r['countries'])-10} (see CSV)" if len(r['countries'])>10 else '')] for r in stances[:9]]))
    models=[r for r in institutions if r['year']==year and r['review_status']=='reviewed' and ('-like' in r['mechanism'] or r['mechanism'].startswith('New body'))]
    if models:third.append(('table',[['Model / proposer','Requested function','Evidence']]+[[r['mechanism']+' / '+r['iso3'],', '.join(r['requested_functions']),str(r['year'])+': '+ ' '.join(r['quote'].split()[:25])+(' …' if len(r['quote'].split())>25 else '')] for r in models[:4]]))
    third.append(('small','Mention, welcome, support, request and commitment are separate. Only explicit UN requests enter UN-request counts. Country positions derive from speech evidence; background proposals and drafts do not enter those counts. Full stance and model tables accompany this report.'))
    fourth=[('title',TITLES[3])]+prose(TITLES[3])
    for r in report['recommendations']:
        fourth.append(('paragraph',r['addressee']+': '+r['action']+' '+r['rationale']+cite(r['evidence_ids']+r['metric_keys'])))
    fourth.append(('small',f"Review limits: automated speech review with sampled second checks and two classification passes; exact-quote checks; no full audio or human verification. "
        f"Unresolved AI speeches: {stats['unresolved_ai_speeches']}; incomplete thematic cases: {stats['incomplete_theme_speeches']}. "
        'The 2025 VCT prepared-only paragraph remains unresolved if its source note persists. Participation rosters are not independently reconciled. Recommendations are analytical proposals.'))
    used={r for s in report['sections'] for p in s['paragraphs'] for r in p['reference_ids']}
    if used:
        fourth.append(('small','Numbered source notes accompany the report in Report_Sources.md; machine-readable evidence and metric registers retain exact source locations and review status.'))
    if any(r.get('status')=='not_verified_at_run' for r in reference_inventory):
        fourth.append(('small','Some official UN web pages could not be retrieved at this run; current institutional status is not fully verified. Draft/background statements are dated and qualified accordingly.'))
    return [first,second,third,fourth]


def source_notes(report,destination):
    import json
    lines=['# Report sources','', 'Numbered citations map to original speech evidence, supplied background and computed statistics. Automated review is not human verification.','']
    for key,item in report.get('citations',{}).items():
        src=item['source'];label=f"{src.get('iso3','')} {src.get('year','')}".strip() or 'Background / computed statistic'
        lines.extend([f"## [{item['number']}] {label}",''])
        if src.get('source_file') or src.get('file'):lines.extend(['Source: '+str(src.get('source_file') or src.get('file')),''])
        if src.get('quote'):lines.extend(['Quotation: '+src['quote'],''])
        lines.extend(['Source locator, review status and/or computed values:','','```json',json.dumps(src,ensure_ascii=False,indent=2),'```',''])
    path=Path(destination)/'Report_Sources.md';path.write_text('\n'.join(lines),encoding='utf-8')
    return str(path)


def render(report_pages,destination,basename,use_word=True):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph,Table,TableStyle,Image,Spacer,PageBreak,SimpleDocTemplate
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from docx import Document
    from docx.shared import Inches,Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    import fitz
    destination=Path(destination);destination.mkdir(parents=True,exist_ok=True)
    font='Helvetica';bold='Helvetica-Bold';windows=Path('C:/Windows/Fonts')
    if (windows/'arial.ttf').exists():
        pdfmetrics.registerFont(TTFont('ReportArial',str(windows/'arial.ttf')))
        pdfmetrics.registerFont(TTFont('ReportArialBold',str(windows/'arialbd.ttf')))
        font='ReportArial';bold='ReportArialBold'
    width,height=A4;content_width=width-88;size=10.5
    styles={kind:ParagraphStyle(kind,fontName=bold if kind=='title' else font,
        fontSize=17 if kind=='title' else 8.5 if kind=='small' else size,
        leading=21 if kind=='title' else 10.5 if kind=='small' else 13,
        keepWithNext=kind=='title',spaceAfter=7,
        textColor=colors.HexColor('#163e54') if kind=='title' else colors.HexColor('#263443')) for kind in ('title','paragraph','small')}
    flows=[]
    for index,section in enumerate(report_pages):
        if index:flows.append(PageBreak())
        for kind,value in section:
            if kind=='image':
                flows.append(Image(value,width=content_width,height=content_width*(2.35/7.2 if 'annual' in value else 2.7/7.2)));flows.append(Spacer(1,7))
            elif kind=='table':
                cells=[[Paragraph(escape(str(c)),styles['small']) for c in row] for row in value]
                widths=[content_width*.46,content_width*.24,content_width*.30] if len(value[0])==3 else [content_width/len(value[0])]*len(value[0])
                table=Table(cells,colWidths=widths,repeatRows=1,hAlign='LEFT')
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e4edf1')),('VALIGN',(0,0),(-1,-1),'TOP'),
                    ('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor('#aac0cb')),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
                flows.extend([table,Spacer(1,7)])
            else:flows.append(Paragraph(escape(value),styles[kind]))
    pdf=destination/(basename+'.pdf')
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont(font,7.5);canvas.setFillColor(colors.HexColor('#637382'))
        canvas.drawString(44,24,'UNGA AI Strategic Review | Automated analysis | Source notes supplied')
        canvas.drawRightString(width-44,24,str(doc.page));canvas.restoreState()
    def build(items):
        SimpleDocTemplate(str(pdf),pagesize=A4,leftMargin=44,rightMargin=44,topMargin=38,bottomMargin=42,
            title='UNGA AI Strategic Review').build(items,onFirstPage=footer,onLaterPages=footer)
    build([flow for flow in flows if not isinstance(flow,PageBreak)])
    with fitz.open(pdf) as check:section_breaks=len(check)<4
    if section_breaks:build(list(flows))
    with fitz.open(pdf) as check:
        if not 4<=len(check)<=6:raise ValueError(f'Report has {len(check)} pages; revise narrative/table selection within 4-6 pages before publication')
    document=Document();section=document.sections[0]
    section.page_width=Inches(width/72);section.page_height=Inches(height/72)
    section.top_margin=Inches(38/72);section.bottom_margin=Inches(42/72)
    section.left_margin=section.right_margin=Inches(44/72)
    normal=document.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(size)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=Pt(13)
    document.styles['Heading 1'].font.size=Pt(17)
    document.styles['Heading 1'].paragraph_format.line_spacing=Pt(21)
    for index,items in enumerate(report_pages):
        if index and section_breaks:document.add_page_break()
        for kind,value in items:
            if kind=='title':document.add_heading(value,1)
            elif kind=='image':
                document.add_picture(value,width=Inches(content_width/72))
                # Exact body line height clips inline drawings in Microsoft Word.
                document.paragraphs[-1].paragraph_format.line_spacing=1.0
            elif kind=='table':
                t=document.add_table(rows=0,cols=len(value[0]));t.style='Light Shading Accent 1'
                for ri,row in enumerate(value):
                    cells=t.add_row().cells
                    if ri==0:t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                    for cell,text in zip(cells,row):
                        cell.text=str(text)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after=Pt(3);paragraph.paragraph_format.line_spacing=Pt(10.5)
                            for run in paragraph.runs:run.font.size=Pt(8.5)
                    t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            else:
                paragraph=document.add_paragraph(value)
                if kind=='small':
                    paragraph.paragraph_format.line_spacing=Pt(10.5)
                    for run in paragraph.runs:run.font.size=Pt(8.5)
    foot=section.footer.paragraphs[0];foot.text='UNGA AI Strategic Review | Automated analysis | '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    docx=destination/(basename+'.docx');document.save(docx)
    word_verified=False;word_note='Word rendering unavailable; flowing PDF used'
    if os.name=='nt' and use_word:
        word_pdf=destination/(basename+'.word-render.pdf')
        try:
            for _ in range(2):
                completed=subprocess.run(['powershell','-NoProfile','-ExecutionPolicy','Bypass','-File',
                    str(Path(__file__).with_name('render_word.ps1')),'-DocxPath',str(docx.resolve()),'-PdfPath',str(word_pdf.resolve())],
                    capture_output=True,timeout=55,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
                if completed.returncode!=0 or not word_pdf.exists():break
                with fitz.open(word_pdf) as word_document:
                    text='\n'.join(p.get_text() for p in word_document);count=len(word_document)
                if count<4 and not section_breaks:
                    section_breaks=True
                    for paragraph in document.paragraphs:
                        if paragraph.style.name=='Heading 1' and paragraph.text!=TITLES[0]:paragraph.paragraph_format.page_break_before=True
                    document.save(docx);continue
                if not 4<=count<=6 or not all(title in text for title in TITLES):
                    raise ValueError('Word-rendered report fails 4-6 page/section requirements; revise content before publication')
                word_pdf.replace(pdf);word_verified=True;word_note='DOCX rendered by installed Microsoft Word; exported PDF is delivered';break
        except (OSError,subprocess.TimeoutExpired):word_note='Word rendering did not complete; flowing PDF used'
    with fitz.open(pdf) as check:
        full_text='\n'.join(p.get_text() for p in check)
        if not all(title in full_text for title in TITLES):raise ValueError('Missing report section in rendered PDF')
        if re.search(r'\bE\d{5}\b|\bannual:\d+',full_text):raise ValueError('Internal citation key leaked into publication')
        for page in check:
            for block in page.get_text('dict')['blocks']:
                if block.get('type')!=0:continue
                for line in block['lines']:
                    for span in line['spans']:
                        x0,y0,x1,y1=span['bbox']
                        if min(x0,y0)<-1 or x1>page.rect.width+1 or y1>page.rect.height+1:raise ValueError('Rendered text exceeds page bounds')
        for i,page in enumerate(check):page.get_pixmap(matrix=fitz.Matrix(1,1)).save(str(destination/f'report_page_{i+1}.png'))
        page_count=len(check)
    return dict(pdf=str(pdf),docx=str(docx),pdf_pages=page_count,pdf_layout_bounds_checked=True,
        docx_has_explicit_page_breaks=section_breaks,sections_can_span_pages=True,docx_render_verified=word_verified,
        docx_render_note=word_note,human_visual_review=False,font_sizes=[size]*page_count)
